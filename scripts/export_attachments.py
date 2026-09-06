"""
Export FastAdmin tables to Excel.

Sources:
  medical      /medical/index          (product catalog)
  attachments  /general/attachment/index

1. Log in in Chrome (solve captcha yourself).
2. Open the table page → DevTools → Network → the index?offset=0 XHR.
3. Copy the FULL Cookie request header into session.txt (gitignored).
4. pip install requests pandas openpyxl
5. python scripts/export_attachments.py --source medical --max-rows 5
   python scripts/export_attachments.py --source all
"""

from __future__ import annotations

import argparse
import json
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

BASE = "https://medicalreport.kingbell.cc"
ADMIN = f"{BASE}/paTGmJjyuV.php"
ROOT = Path(__file__).resolve().parents[1]
COOKIE_FILE = ROOT / "session.txt"
OUT_DIR = ROOT / "exports"
FILES_DIR = OUT_DIR / "files"

TITLE_FILL = "0F766E"
HEADER_FILL = "134E4A"
ALT_FILL = "F0FDFA"
TEXT = "1C1917"
SUBTITLE = "134E4A"


def absolute_url(path: object) -> str:
    if path is None:
        return ""
    text = str(path).strip()
    if not text:
        return ""
    if text.startswith("http://") or text.startswith("https://"):
        return text
    if not text.startswith("/"):
        text = "/" + text
    return f"{BASE}{text}"


def split_paths(value: object) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    return [part.strip() for part in re.split(r"[,;\n]+", text) if part.strip()]


SOURCES: dict[str, dict[str, Any]] = {
    "medical": {
        "title": "Medical equipment catalog",
        "sheet": "Catalog",
        "path": "/medical/index",
        "sort": "id",
        "order": "asc",
        "extra_params": {"addtabs": "1"},
        "columns": [
            ("id", "ID"),
            ("ref", "Ref"),
            ("product_name", "Product"),
            ("chinese_name", "Chinese name"),
            ("brand", "Brand"),
            ("model", "Model"),
            ("qty", "Qty"),
            ("price", "Price"),
            ("department", "Department"),
            ("status_text", "Status"),
            ("memo", "Memo"),
            ("file_urls", "Datasheet URLs"),
            ("ce_url", "CE certificate"),
        ],
        "hyperlink_columns": ("Datasheet URLs", "CE certificate"),
        "widths": {
            "ID": 10,
            "Ref": 8,
            "Product": 36,
            "Chinese name": 28,
            "Brand": 14,
            "Model": 18,
            "Qty": 8,
            "Price": 14,
            "Department": 16,
            "Status": 12,
            "Memo": 24,
            "Datasheet URLs": 52,
            "CE certificate": 52,
        },
    },
    "attachments": {
        "title": "Attachment files",
        "sheet": "Attachments",
        "path": "/general/attachment/index",
        "sort": "id",
        "order": "desc",
        "extra_params": {},
        "columns": [
            ("id", "ID"),
            ("filename", "Filename"),
            ("chinese_name", "Chinese name"),
            ("model", "Model"),
            ("imagetype", "Type"),
            ("filesize_kb", "Size (KB)"),
            ("mimetype", "MIME"),
            ("category", "Category"),
            ("createtime_iso", "Uploaded (UTC)"),
            ("file_url", "File URL"),
        ],
        "hyperlink_columns": ("File URL",),
        "widths": {
            "ID": 10,
            "Filename": 42,
            "Chinese name": 22,
            "Model": 28,
            "Type": 10,
            "Size (KB)": 12,
            "MIME": 22,
            "Category": 14,
            "Uploaded (UTC)": 20,
            "File URL": 56,
        },
    },
}


def load_cookie(path: Path) -> str:
    if not path.exists():
        raise SystemExit(
            f"Missing {path}. Copy the full Cookie header from DevTools into that file."
        )
    text = path.read_text(encoding="utf-8").strip()
    text = re.sub(r"^Cookie:\s*", "", text, flags=re.I)
    if not text:
        raise SystemExit(f"{path} is empty.")
    parts = [p.strip() for p in text.split(";") if p.strip() and "=" in p]
    print(f"Loaded {len(parts)} cookie(s) from {path.name}")
    if len(parts) < 3:
        print(
            "Warning: FastAdmin usually sends more than PHPSESSID. "
            "Copy the entire Cookie header from the table XHR if login fails."
        )
    return text


def make_session(cookie: str) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "X-Requested-With": "XMLHttpRequest",
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            "Cookie": cookie,
        }
    )
    return session


def unix_to_iso(value: object) -> str | None:
    try:
        n = int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    if n <= 0:
        return None
    return datetime.fromtimestamp(n, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def parse_payload(response: requests.Response) -> tuple[int, list[dict]]:
    if "text/html" in response.headers.get("Content-Type", "") or response.text.lstrip().startswith(
        "<"
    ):
        raise SystemExit(
            "Got an HTML login page. Session expired — log in again and update session.txt."
        )
    try:
        payload = response.json()
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Response was not JSON: {exc}") from exc

    if isinstance(payload, dict) and payload.get("code") not in (None, 1, "1"):
        raise SystemExit(
            f"API error code={payload.get('code')} msg={payload.get('msg')!r}"
        )

    if isinstance(payload, dict) and "rows" in payload:
        rows = payload.get("rows") or []
        return int(payload.get("total") or len(rows)), rows

    data = payload.get("data") if isinstance(payload, dict) else None
    if isinstance(data, dict) and "rows" in data:
        rows = data.get("rows") or []
        return int(data.get("total") or len(rows)), rows
    if isinstance(data, list):
        return len(data), data

    raise SystemExit(f"Unexpected JSON shape: {list(payload)[:8]!r} data={data!r}")


def fetch_source(
    session: requests.Session,
    source: dict[str, Any],
    limit: int,
    delay: float,
    max_rows: int | None,
) -> tuple[list[dict], int]:
    url = f"{ADMIN}{source['path']}"
    session.headers["Referer"] = f"{url}?addtabs=1"
    rows: list[dict] = []
    offset = 0
    total = None
    while True:
        page_limit = limit
        if max_rows is not None:
            page_limit = min(limit, max_rows - len(rows))
            if page_limit <= 0:
                break
        params = {
            "sort": source["sort"],
            "order": source["order"],
            "offset": offset,
            "limit": page_limit,
            "filter": "{}",
            "op": "{}",
            "_": int(time.time() * 1000),
            **source.get("extra_params", {}),
        }
        response = session.get(url, params=params, timeout=60)
        response.raise_for_status()
        page_total, batch = parse_payload(response)
        if total is None:
            total = page_total
            print(f"  {source['sheet']}: server reports {total} rows")
        if not batch:
            break
        rows.extend(batch)
        print(f"  {source['sheet']}: fetched {len(rows)} / {total}")
        if max_rows is not None and len(rows) >= max_rows:
            rows = rows[:max_rows]
            break
        if len(rows) >= total or len(batch) < page_limit:
            break
        offset += page_limit
        time.sleep(delay)
    return rows, int(total or len(rows))


def frame_medical(rows: list[dict]) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    frame["qty"] = pd.to_numeric(frame.get("qty"), errors="coerce")
    frame["price"] = pd.to_numeric(frame.get("price"), errors="coerce")
    frame["file_urls"] = frame.get("download_files", pd.Series(dtype=str)).map(
        lambda value: "\n".join(absolute_url(p) for p in split_paths(value))
    )
    frame["ce_url"] = frame.get("ce", pd.Series(dtype=str)).map(absolute_url)
    keep = [key for key, _ in SOURCES["medical"]["columns"] if key in frame.columns]
    return frame[keep]


def frame_attachments(rows: list[dict]) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    frame["filesize_kb"] = (
        pd.to_numeric(frame.get("filesize"), errors="coerce") / 1024
    ).round(1)
    if "createtime" in frame.columns:
        frame["createtime_iso"] = frame["createtime"].map(unix_to_iso)
    rel = frame.get("fullurl", frame.get("url"))
    frame["file_url"] = rel.map(absolute_url)
    keep = [key for key, _ in SOURCES["attachments"]["columns"] if key in frame.columns]
    return frame[keep]


FRAME_BUILDERS = {
    "medical": frame_medical,
    "attachments": frame_attachments,
}


def style_sheet(
    sheet,
    pretty: pd.DataFrame,
    title: str,
    total_on_server: int,
    hyperlink_columns: tuple[str, ...],
    widths: dict[str, int],
) -> None:
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=10, color=SUBTITLE)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color=TEXT)
    link_font = Font(name="Calibri", size=10, color=TITLE_FILL, underline="single")
    title_fill = PatternFill("solid", fgColor=TITLE_FILL)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    alt_fill = PatternFill("solid", fgColor=ALT_FILL)
    thin = Border(
        left=Side(style="thin", color="D6D3D1"),
        right=Side(style="thin", color="D6D3D1"),
        top=Side(style="thin", color="D6D3D1"),
        bottom=Side(style="thin", color="D6D3D1"),
    )

    col_count = max(len(pretty.columns), 1)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    sheet["A1"] = title
    sheet["A1"].font = title_font
    sheet["A1"].fill = title_fill
    sheet["A1"].alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 28
    sheet["A2"] = (
        f"{len(pretty)} rows in this sheet · {total_on_server} on server · "
        f"exported {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    sheet["A2"].font = subtitle_font
    sheet["A2"].alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[2].height = 20

    for r_idx, row in enumerate(
        dataframe_to_rows(pretty, index=False, header=True), start=3
    ):
        for c_idx, value in enumerate(row, start=1):
            cell = sheet.cell(r_idx, c_idx, value)
            cell.font = header_font if r_idx == 3 else body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin
            if r_idx == 3:
                cell.fill = header_fill
            elif r_idx % 2 == 0:
                cell.fill = alt_fill

    first_data = 4
    last_data = 3 + len(pretty)
    for col_name in hyperlink_columns:
        if col_name not in pretty.columns:
            continue
        col_idx = list(pretty.columns).index(col_name) + 1
        for r in range(first_data, last_data + 1):
            cell = sheet.cell(r, col_idx)
            raw = str(cell.value or "").strip()
            first = raw.split("\n")[0].strip()
            if first.startswith("http"):
                cell.hyperlink = first
                cell.font = link_font

    for idx, name in enumerate(pretty.columns, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = widths.get(name, 16)

    if pretty.columns.size:
        sheet.auto_filter.ref = f"A3:{get_column_letter(col_count)}{max(last_data, 3)}"
    sheet.freeze_panes = "A4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:3"
    sheet.sheet_view.showGridLines = False


def write_workbook(sheets: list[dict[str, Any]], path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    for item in sheets:
        source = item["source"]
        pretty = item["frame"].rename(columns=dict(source["columns"]))
        sheet = book.create_sheet(source["sheet"])
        style_sheet(
            sheet,
            pretty,
            title=source["title"],
            total_on_server=item["total"],
            hyperlink_columns=source["hyperlink_columns"],
            widths=source["widths"],
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def download_files(session: requests.Session, rows: list[dict], delay: float) -> None:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    for i, row in enumerate(rows, start=1):
        paths = split_paths(row.get("download_files")) + split_paths(row.get("ce"))
        paths += split_paths(row.get("fullurl") or row.get("url"))
        for rel in paths:
            url = absolute_url(rel)
            name = re.sub(r"[^\w.\-]+", "_", Path(rel).name or "file")
            dest = FILES_DIR / f"{row.get('id')}_{name}"
            if dest.exists():
                continue
            response = session.get(url, timeout=120)
            if response.status_code != 200:
                print(f"Skip {url} ({response.status_code})")
                continue
            dest.write_bytes(response.content)
        if i % 25 == 0:
            print(f"Downloaded files for {i} / {len(rows)} rows")
        time.sleep(delay)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export admin tables to Excel")
    parser.add_argument("--cookie-file", type=Path, default=COOKIE_FILE)
    parser.add_argument(
        "--source",
        choices=["medical", "attachments", "all"],
        default="all",
    )
    parser.add_argument("--limit", type=int, default=100)
    parser.add_argument("--max-rows", type=int, default=None)
    parser.add_argument("--delay", type=float, default=0.2)
    parser.add_argument("--files", action="store_true")
    parser.add_argument("--out", type=Path, default=OUT_DIR / "catalog.xlsx")
    args = parser.parse_args()

    names = list(SOURCES) if args.source == "all" else [args.source]
    cookie = load_cookie(args.cookie_file)
    session = make_session(cookie)

    sheets: list[dict[str, Any]] = []
    all_rows: list[dict] = []
    for name in names:
        source = SOURCES[name]
        print(f"Exporting {name}…")
        rows, total = fetch_source(
            session, source, args.limit, args.delay, args.max_rows
        )
        if not rows:
            print(f"  {source['sheet']}: no rows")
            continue
        frame = FRAME_BUILDERS[name](rows)
        sheets.append({"source": source, "frame": frame, "total": total})
        all_rows.extend(rows)

    if not sheets:
        raise SystemExit("No rows returned.")

    write_workbook(sheets, args.out)
    print(f"Wrote {args.out}")

    if args.files:
        download_files(session, all_rows, args.delay)
        print(f"Files saved under {FILES_DIR}")


if __name__ == "__main__":
    main()
