"""
Duplicate exports/catalog.xlsx and enrich the copy.

Does not modify the original workbook.
Adds Translation + Local files columns, downloads PDFs/PNGs into exports/files/,
and writes clickable local hyperlinks.

  python scripts/enrich_catalog.py
  python scripts/enrich_catalog.py --skip-download
  python scripts/enrich_catalog.py --max-files 5
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_attachments import COOKIE_FILE, load_cookie, make_session, split_paths

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "exports"
SOURCE = OUT_DIR / "catalog.xlsx"
DEST = OUT_DIR / "catalog-enriched.xlsx"
FILES_DIR = OUT_DIR / "files"
CACHE_PATH = OUT_DIR / "translation-cache.json"
ZIP_PATH = OUT_DIR / "catalog-enriched.zip"

CJK = re.compile(r"[\u4e00-\u9fff]")
TITLE_FILL = "0F766E"
HEADER_FILL = "134E4A"
ALT_FILL = "F0FDFA"


def looks_cjk(text: object) -> bool:
    return bool(text) and bool(CJK.search(str(text)))


def looks_english(text: object) -> bool:
    s = str(text or "").strip()
    if not s:
        return False
    return not looks_cjk(s)


def parse_urls(value: object) -> list[str]:
    urls: list[str] = []
    for part in split_paths(value):
        if part.startswith("http://") or part.startswith("https://"):
            urls.append(part)
    return urls


def file_for_url(url: str) -> str:
    digest = hashlib.sha1(url.encode("utf-8")).hexdigest()[:10]
    name = Path(url.split("?")[0]).name or "file"
    name = re.sub(r"[^\w.\-]+", "_", name)
    return f"{digest}_{name}"


def load_cache() -> dict[str, str]:
    if CACHE_PATH.exists():
        return json.loads(CACHE_PATH.read_text(encoding="utf-8"))
    return {}


def save_cache(cache: dict[str, str]) -> None:
    CACHE_PATH.write_text(
        json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def translate_missing(texts: list[str], cache: dict[str, str]) -> dict[str, str]:
    pending = [t for t in texts if t not in cache]
    if not pending:
        return cache
    try:
        from deep_translator import GoogleTranslator
    except ImportError:
        print("deep-translator not installed; leaving unmatched Chinese untranslated")
        return cache

    translator = GoogleTranslator(source="zh-CN", target="en")
    print(f"Translating {len(pending)} unique Chinese strings…")
    for i, text in enumerate(pending, start=1):
        try:
            cache[text] = translator.translate(text) or text
        except Exception as exc:
            print(f"  skip translation ({exc})")
            cache[text] = text
        if i % 25 == 0:
            save_cache(cache)
            print(f"  translated {i} / {len(pending)}")
        time.sleep(0.15)
    save_cache(cache)
    return cache


def catalog_translation_map(catalog: pd.DataFrame) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for _, row in catalog.iterrows():
        zh = str(row.get("Chinese name") or "").strip()
        en = str(row.get("Product") or "").strip()
        if zh and looks_english(en):
            mapping[zh] = en
    return mapping


def translation_for(zh: object, product: object, mapping: dict[str, str], cache: dict[str, str]) -> str:
    product_s = str(product or "").strip()
    if looks_english(product_s):
        return product_s
    zh_s = str(zh or "").strip()
    if not zh_s:
        return ""
    if not looks_cjk(zh_s):
        return zh_s
    return mapping.get(zh_s) or cache.get(zh_s) or ""


def read_sheet(path: Path, name: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=name, header=2)


def local_files_cell(urls: list[str]) -> str:
    return "\n".join(f"files/{file_for_url(url)}" for url in urls)


def download_url(session: requests.Session, url: str, dest: Path) -> bool:
    if dest.exists() and dest.stat().st_size > 0:
        return True
    tmp = dest.with_name(f"{dest.name}.{uuid.uuid4().hex[:8]}.part")
    try:
        response = session.get(url, timeout=(8, 25), stream=True)
        content_type = response.headers.get("Content-Type", "")
        if response.status_code != 200 or "text/html" in content_type.lower():
            print(f"  skip {dest.name} ({response.status_code} {content_type.split(';')[0]})")
            return False
        dest.parent.mkdir(parents=True, exist_ok=True)
        with tmp.open("wb") as handle:
            for chunk in response.iter_content(chunk_size=64 * 1024):
                if chunk:
                    handle.write(chunk)
        if not tmp.exists() or tmp.stat().st_size == 0:
            if tmp.exists():
                tmp.unlink()
            print(f"  skip {dest.name} (empty)")
            return False
        try:
            if dest.exists():
                dest.unlink()
        except OSError:
            pass
        shutil.copyfile(tmp, dest)
        try:
            tmp.unlink()
        except OSError:
            pass
        return dest.exists() and dest.stat().st_size > 0
    except Exception as exc:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
        print(f"  skip {dest.name} ({exc})")
        return False


def collect_downloads(catalog: pd.DataFrame, attachments: pd.DataFrame) -> list[tuple[str, Path]]:
    jobs: list[tuple[str, Path]] = []
    seen: set[str] = set()
    for urls in (
        [parse_urls(row.get("Datasheet URLs")) + parse_urls(row.get("CE certificate")) for _, row in catalog.iterrows()]
        + [parse_urls(row.get("File URL")) for _, row in attachments.iterrows()]
    ):
        for url in urls:
            if url in seen:
                continue
            seen.add(url)
            jobs.append((url, FILES_DIR / file_for_url(url)))
    return jobs


def style_sheet(sheet, pretty: pd.DataFrame, title: str, note: str, hyperlink_cols: tuple[str, ...]) -> None:
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=10, color="134E4A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="1C1917")
    link_font = Font(name="Calibri", size=10, color="0F766E", underline="single")
    title_fill = PatternFill("solid", fgColor=TITLE_FILL)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    alt_fill = PatternFill("solid", fgColor=ALT_FILL)
    thin = Border(
        left=Side(style="thin", color="D6D3D1"),
        right=Side(style="thin", color="D6D3D1"),
        top=Side(style="thin", color="D6D3D1"),
        bottom=Side(style="thin", color="D6D3D1"),
    )
    col_count = max(len(pretty.columns), 1)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    sheet["A1"] = title
    sheet["A1"].font = title_font
    sheet["A1"].fill = title_fill
    sheet["A1"].alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 28
    sheet["A2"] = note
    sheet["A2"].font = subtitle_font
    sheet["A2"].alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[2].height = 20

    for r_idx, row in enumerate(dataframe_to_rows(pretty, index=False, header=True), start=3):
        for c_idx, value in enumerate(row, start=1):
            cell = sheet.cell(r_idx, c_idx, value)
            cell.font = header_font if r_idx == 3 else body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin
            if r_idx == 3:
                cell.fill = header_fill
            elif r_idx % 2 == 0:
                cell.fill = alt_fill

    last = 3 + len(pretty)
    for col_name in hyperlink_cols:
        if col_name not in pretty.columns:
            continue
        col_idx = list(pretty.columns).index(col_name) + 1
        for r in range(4, last + 1):
            cell = sheet.cell(r, col_idx)
            raw = str(cell.value or "").strip()
            first = raw.split("\n")[0].strip()
            if first.startswith("http") or first.startswith("files/"):
                cell.hyperlink = first
                cell.font = link_font

    widths = {
        "Translation": 36,
        "Local files": 42,
        "Datasheet URLs": 48,
        "CE certificate": 48,
        "File URL": 48,
        "Product": 36,
        "Chinese name": 28,
        "Filename": 42,
    }
    for idx, name in enumerate(pretty.columns, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)
    sheet.auto_filter.ref = f"A3:{get_column_letter(col_count)}{max(last, 3)}"
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False


def insert_after(df: pd.DataFrame, after: str, name: str, values: pd.Series) -> pd.DataFrame:
    cols = list(df.columns)
    if name in cols:
        df[name] = values
        return df
    idx = cols.index(after) + 1 if after in cols else len(cols)
    df.insert(idx, name, values)
    return df


def write_workbook(catalog: pd.DataFrame, attachments: pd.DataFrame, dest: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    for name, frame, title in (
        (
            "Catalog",
            catalog,
            "Medical equipment catalog (enriched copy)",
        ),
        (
            "Attachments",
            attachments,
            "Attachment files (enriched copy)",
        ),
    ):
        sheet = book.create_sheet(name)
        style_sheet(
            sheet,
            frame,
            title=title,
            note=(
                f"{len(frame)} rows · original {SOURCE.name} left unchanged · "
                f"local files in the files folder next to this workbook"
            ),
            hyperlink_cols=(
                "Datasheet URLs",
                "CE certificate",
                "File URL",
                "Local files",
            ),
        )
    dest.parent.mkdir(parents=True, exist_ok=True)
    book.save(dest)


def zip_pack(dest: Path) -> None:
    print(f"Zipping {ZIP_PATH.name}…")
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.write(dest, dest.name)
        if FILES_DIR.exists():
            for path in FILES_DIR.rglob("*"):
                if path.is_file():
                    zf.write(path, Path("files") / path.name)
    print(f"Wrote {ZIP_PATH} ({ZIP_PATH.stat().st_size / (1024 * 1024):.1f} MB)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--out", type=Path, default=DEST)
    parser.add_argument("--skip-download", action="store_true")
    parser.add_argument("--skip-zip", action="store_true")
    parser.add_argument("--max-files", type=int, default=None)
    parser.add_argument("--delay", type=float, default=0.12)
    args = parser.parse_args()

    if not args.source.exists():
        raise SystemExit(f"Missing original workbook {args.source}")

    print(f"Leaving original untouched: {args.source}")
    catalog = read_sheet(args.source, "Catalog")
    attachments = read_sheet(args.source, "Attachments")

    mapping = catalog_translation_map(catalog)
    cache = load_cache()
    needed: list[str] = []
    for zh, product in zip(
        list(catalog.get("Chinese name", pd.Series(dtype=str)))
        + list(attachments.get("Chinese name", pd.Series(dtype=str))),
        list(catalog.get("Product", pd.Series(dtype=str)))
        + [""] * len(attachments),
    ):
        zh_s = "" if pd.isna(zh) else str(zh).strip()
        product_s = "" if pd.isna(product) else str(product).strip()
        if looks_cjk(zh_s) and not looks_english(product_s) and zh_s not in mapping and zh_s not in cache:
            needed.append(zh_s)
    cache = translate_missing(sorted(set(needed)), cache)

    catalog = insert_after(
        catalog,
        "Chinese name",
        "Translation",
        [
            translation_for(zh, product, mapping, cache)
            for zh, product in zip(
                catalog.get("Chinese name", pd.Series(dtype=str)).fillna(""),
                catalog.get("Product", pd.Series(dtype=str)).fillna(""),
            )
        ],
    )
    attachments = insert_after(
        attachments,
        "Chinese name",
        "Translation",
        [
            translation_for(zh, "", mapping, cache)
            for zh in attachments.get("Chinese name", pd.Series(dtype=str)).fillna("")
        ],
    )

    catalog["Local files"] = [
        local_files_cell(
            parse_urls(row.get("Datasheet URLs")) + parse_urls(row.get("CE certificate")),
        )
        for _, row in catalog.iterrows()
    ]
    attachments["Local files"] = [
        local_files_cell(parse_urls(row.get("File URL")))
        for _, row in attachments.iterrows()
    ]

    catalog = catalog.fillna("")
    attachments = attachments.fillna("")

    write_workbook(catalog, attachments, args.out)
    print(f"Wrote enriched copy → {args.out}")
    print(f"Original still at → {args.source}")

    jobs = collect_downloads(catalog, attachments)
    if args.max_files:
        jobs = jobs[: args.max_files]
    print(f"{len(jobs)} unique files to download")

    if not args.skip_download:
        cookie = load_cookie(COOKIE_FILE)
        pending = [
            (url, dest)
            for url, dest in jobs
            if not (dest.exists() and dest.stat().st_size > 0)
        ]
        print(f"{len(pending)} remaining after skip-existing")
        ok = len(jobs) - len(pending)

        def worker(item: tuple[str, Path]) -> bool:
            url, dest = item
            session = make_session(cookie)
            return download_url(session, url, dest)

        attempted = 0
        with ThreadPoolExecutor(max_workers=3) as pool:
            futures = [pool.submit(worker, item) for item in pending]
            for future in as_completed(futures):
                attempted += 1
                if future.result():
                    ok += 1
                if attempted % 25 == 0 or attempted == len(pending):
                    print(f"  downloaded {ok} ok · {attempted} / {len(pending)} new attempts")
        print(f"Downloaded {ok} / {len(jobs)} files → {FILES_DIR}")

    if not args.skip_zip:
        zip_pack(args.out)


if __name__ == "__main__":
    main()
